from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock

from openpyxl import Workbook
from docx import Document

from app.core.config import get_settings
from app.core.database import available, connection


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _new_id() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S%f")


@dataclass
class GeneratedFileRecord:
    id: str
    name: str
    type: str
    size: int
    chat_id: str | None
    created_at: datetime
    stored_as: str


class GeneratedFileStore:
    def __init__(self) -> None:
        s = get_settings()
        self._dir = s.generated_dir
        self._dir.mkdir(parents=True, exist_ok=True)
        self._path = s.storage_dir / "generated_files.json"
        self._lock = Lock()

    def _file_path(self, stored_as: str) -> Path:
        return self._dir / stored_as

    def list_all(self) -> list[GeneratedFileRecord]:
        if available():
            with connection() as conn:
                rows = conn.execute(
                    "SELECT id, name, type, size, chat_id, created_at, stored_as FROM generated_files ORDER BY created_at DESC"
                ).fetchall()
            return [GeneratedFileRecord(*r) for r in rows]
        if not self._path.exists():
            return []
        import json

        items = json.loads(self._path.read_text(encoding="utf-8"))
        return [GeneratedFileRecord(**i) for i in items]

    def get(self, file_id: str) -> GeneratedFileRecord | None:
        if available():
            with connection() as conn:
                row = conn.execute(
                    "SELECT id, name, type, size, chat_id, created_at, stored_as FROM generated_files WHERE id = %s",
                    (file_id,),
                ).fetchone()
            return GeneratedFileRecord(*row) if row else None
        return next((f for f in self.list_all() if f.id == file_id), None)

    def create_from_chat(self, chat_id: str, prompt: str, answer: str) -> GeneratedFileRecord | None:
        prompt_l = prompt.lower()
        if any(k in prompt_l for k in ("excel", "xlsx", "spreadsheet", "sheet")):
            return self._create_xlsx(chat_id, prompt, answer)
        if any(k in prompt_l for k in ("word", "docx", "document")):
            return self._create_docx(chat_id, prompt, answer)
        return self._create_text(chat_id, prompt, answer)

    def _persist(self, record: GeneratedFileRecord) -> None:
        if available():
            with connection() as conn:
                conn.execute(
                    "INSERT INTO generated_files (id, name, type, size, chat_id, created_at, stored_as) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                    (record.id, record.name, record.type, record.size, record.chat_id, record.created_at, record.stored_as),
                )
            return
        import json

        with self._lock:
            items = self.list_all()
            items.append(record)
            self._path.write_text(json.dumps([r.__dict__ for r in items], default=str, ensure_ascii=False, indent=2), encoding="utf-8")

    def _create_text(self, chat_id: str, prompt: str, answer: str) -> GeneratedFileRecord:
        file_id = _new_id()
        name = "generated_summary.txt"
        stored_as = f"{file_id}.txt"
        path = self._file_path(stored_as)
        content = f"Prompt:\n{prompt}\n\nAnswer:\n{answer}\n"
        path.write_text(content, encoding="utf-8")
        record = GeneratedFileRecord(id=file_id, name=name, type="txt", size=path.stat().st_size, chat_id=chat_id, created_at=_now(), stored_as=stored_as)
        self._persist(record)
        return record

    def _create_docx(self, chat_id: str, prompt: str, answer: str) -> GeneratedFileRecord:
        file_id = _new_id()
        name = "generated_report.docx"
        stored_as = f"{file_id}.docx"
        path = self._file_path(stored_as)
        doc = Document()
        doc.add_heading("Generated Report", level=1)
        doc.add_paragraph(f"Prompt: {prompt}")
        doc.add_paragraph(answer)
        doc.save(path)
        record = GeneratedFileRecord(id=file_id, name=name, type="docx", size=path.stat().st_size, chat_id=chat_id, created_at=_now(), stored_as=stored_as)
        self._persist(record)
        return record

    def _create_xlsx(self, chat_id: str, prompt: str, answer: str) -> GeneratedFileRecord:
        file_id = _new_id()
        name = "generated_checklist.xlsx"
        stored_as = f"{file_id}.xlsx"
        path = self._file_path(stored_as)
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.append(["Prompt", prompt])
        ws.append([])
        ws.append(["Answer"])
        for line in answer.splitlines():
            ws.append([line])
        wb.save(path)
        record = GeneratedFileRecord(id=file_id, name=name, type="xlsx", size=path.stat().st_size, chat_id=chat_id, created_at=_now(), stored_as=stored_as)
        self._persist(record)
        return record

    def delete(self, file_id: str) -> bool:
        record = self.get(file_id)
        if not record:
            return False
        if available():
            with connection() as conn:
                conn.execute("DELETE FROM generated_files WHERE id = %s", (file_id,))
        else:
            import json

            items = [r for r in self.list_all() if r.id != file_id]
            self._path.write_text(json.dumps([r.__dict__ for r in items], default=str, ensure_ascii=False, indent=2), encoding="utf-8")
        self._file_path(record.stored_as).unlink(missing_ok=True)
        return True

    def path_for(self, file_id: str) -> Path | None:
        record = self.get(file_id)
        return self._file_path(record.stored_as) if record else None


generated_file_store = GeneratedFileStore()
